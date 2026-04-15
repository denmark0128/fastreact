import csv
from datetime import date
from io import BytesIO, StringIO

from fastapi import HTTPException, status
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.attendance.models import AttendanceRecord
from app.attendance.schemas import AttendanceRecordCreate, AttendanceRecordResponse, AttendanceRecordUpdate
from app.employees.models import Employee


def _normalize_time(value: str | None) -> str | None:
    if value is None:
        return None

    stripped = str(value).strip()
    if not stripped:
        return None

    parts = stripped.split(":")
    if len(parts) != 2:
        raise ValueError("Time must be in HH:MM format")

    h, m = parts
    if not (h.isdigit() and m.isdigit()):
        raise ValueError("Time must be in HH:MM format")

    hour = int(h)
    minute = int(m)
    if not (0 <= hour <= 23 and 0 <= minute <= 59):
        raise ValueError("Invalid time value")

    return f"{hour:02d}:{minute:02d}"


def _validate_iso_date(value: str) -> str:
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise ValueError("date must be in YYYY-MM-DD format") from exc


def _get_employee(db: Session, employee_id: int) -> Employee:
    employee = db.query(Employee).filter(Employee.id == employee_id).first()
    if not employee:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Employee not found")
    return employee


def _serialize_record(record: AttendanceRecord, employee: Employee | None = None) -> AttendanceRecordResponse:
    return AttendanceRecordResponse(
        id=record.id,
        employee_id=record.employee_id,
        employee_name=employee.profile_name if employee else None,
        employee_code=employee.employee_code if employee else None,
        date=record.date,
        am_in=record.am_in,
        am_out=record.am_out,
        pm_in=record.pm_in,
        pm_out=record.pm_out,
        note=record.note,
        created_at=record.created_at,
        updated_at=record.updated_at,
    )


def list_attendance_records(
    db: Session,
    *,
    employee_id: int | None,
    date_from: str | None,
    date_to: str | None,
    q: str | None,
    skip: int,
    limit: int,
) -> tuple[list[AttendanceRecordResponse], int]:
    query = db.query(AttendanceRecord, Employee).join(Employee, Employee.id == AttendanceRecord.employee_id)

    if employee_id is not None:
        query = query.filter(AttendanceRecord.employee_id == employee_id)
    if date_from:
        query = query.filter(AttendanceRecord.date >= _validate_iso_date(date_from))
    if date_to:
        query = query.filter(AttendanceRecord.date <= _validate_iso_date(date_to))
    if q and q.strip():
        keyword = f"%{q.strip()}%"
        query = query.filter(
            or_(
                Employee.profile_name.ilike(keyword),
                Employee.employee_code.ilike(keyword),
                AttendanceRecord.date.ilike(keyword),
                AttendanceRecord.note.ilike(keyword),
                AttendanceRecord.am_in.ilike(keyword),
                AttendanceRecord.pm_out.ilike(keyword),
            )
        )

    total = query.count()
    rows = query.order_by(AttendanceRecord.date.desc(), AttendanceRecord.id.desc()).offset(skip).limit(limit).all()

    items = [_serialize_record(record, employee) for record, employee in rows]
    return items, total


def create_attendance_record(db: Session, payload: AttendanceRecordCreate) -> AttendanceRecordResponse:
    employee = _get_employee(db, payload.employee_id)

    date_value = _validate_iso_date(payload.date)
    existing = (
        db.query(AttendanceRecord)
        .filter(AttendanceRecord.employee_id == payload.employee_id, AttendanceRecord.date == date_value)
        .first()
    )
    if existing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Attendance record already exists for this employee/date",
        )

    record = AttendanceRecord(
        employee_id=payload.employee_id,
        date=date_value,
        am_in=_normalize_time(payload.am_in),
        am_out=_normalize_time(payload.am_out),
        pm_in=_normalize_time(payload.pm_in),
        pm_out=_normalize_time(payload.pm_out),
        note=payload.note,
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return _serialize_record(record, employee)


def update_attendance_record(db: Session, record_id: int, payload: AttendanceRecordUpdate) -> AttendanceRecordResponse:
    record = db.query(AttendanceRecord).filter(AttendanceRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Attendance record not found")

    if payload.am_in is not None:
        record.am_in = _normalize_time(payload.am_in)
    if payload.am_out is not None:
        record.am_out = _normalize_time(payload.am_out)
    if payload.pm_in is not None:
        record.pm_in = _normalize_time(payload.pm_in)
    if payload.pm_out is not None:
        record.pm_out = _normalize_time(payload.pm_out)
    if payload.note is not None:
        record.note = payload.note

    db.commit()
    db.refresh(record)

    employee = db.query(Employee).filter(Employee.id == record.employee_id).first()
    return _serialize_record(record, employee)


def _canonical_header(header: str) -> str:
    return header.strip().lower().replace(" ", "_")


def _resolve_column_name(column: str) -> str:
    aliases = {
        "employee_id": "employee_id",
        "emp_id": "employee_id",
        "employee_code": "employee_code",
        "emp_code": "employee_code",
        "date": "date",
        "attendance_date": "date",
        "am_in": "am_in",
        "before_noon_in": "am_in",
        "morning_in": "am_in",
        "am_out": "am_out",
        "before_noon_out": "am_out",
        "morning_out": "am_out",
        "pm_in": "pm_in",
        "after_noon_in": "pm_in",
        "afternoon_in": "pm_in",
        "pm_out": "pm_out",
        "after_noon_out": "pm_out",
        "afternoon_out": "pm_out",
        "note": "note",
        "notes": "note",
    }
    return aliases.get(_canonical_header(column), _canonical_header(column))


def parse_csv_rows(file_bytes: bytes) -> list[dict[str, str]]:
    content = file_bytes.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(content))
    rows: list[dict[str, str]] = []

    for row in reader:
        if not row:
            continue
        mapped: dict[str, str] = {}
        for key, value in row.items():
            if key is None:
                continue
            mapped[_resolve_column_name(key)] = "" if value is None else str(value).strip()
        if any(value for value in mapped.values()):
            rows.append(mapped)

    return rows


def parse_excel_rows(file_bytes: bytes) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Excel support requires openpyxl package",
        ) from exc

    workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True)
    sheet = workbook.active

    raw_headers = [str(cell).strip() if cell is not None else "" for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])]
    headers = [_resolve_column_name(header) for header in raw_headers]

    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        mapped: dict[str, str] = {}
        for index, value in enumerate(values):
            header = headers[index] if index < len(headers) else ""
            if not header:
                continue
            mapped[header] = "" if value is None else str(value).strip()
        if any(value for value in mapped.values()):
            rows.append(mapped)

    return rows


def import_attendance_records(db: Session, rows: list[dict[str, str]]) -> dict[str, object]:
    imported = 0
    updated = 0
    skipped = 0
    errors: list[str] = []

    employee_by_code = {
        employee.employee_code: employee
        for employee in db.query(Employee).all()
        if employee.employee_code
    }

    for index, row in enumerate(rows, start=2):
        try:
            employee_id_value = row.get("employee_id", "").strip()
            employee_code_value = row.get("employee_code", "").strip()
            date_value = row.get("date", "").strip()

            if not date_value:
                skipped += 1
                errors.append(f"Row {index}: missing date")
                continue

            normalized_date = _validate_iso_date(date_value)

            employee: Employee | None = None
            if employee_id_value:
                if not employee_id_value.isdigit():
                    raise ValueError("employee_id must be numeric")
                employee = db.query(Employee).filter(Employee.id == int(employee_id_value)).first()
            elif employee_code_value:
                employee = employee_by_code.get(employee_code_value)

            if not employee:
                skipped += 1
                errors.append(f"Row {index}: employee not found")
                continue

            existing = (
                db.query(AttendanceRecord)
                .filter(AttendanceRecord.employee_id == employee.id, AttendanceRecord.date == normalized_date)
                .first()
            )

            am_in = _normalize_time(row.get("am_in"))
            am_out = _normalize_time(row.get("am_out"))
            pm_in = _normalize_time(row.get("pm_in"))
            pm_out = _normalize_time(row.get("pm_out"))
            note = row.get("note") or None

            if existing:
                existing.am_in = am_in
                existing.am_out = am_out
                existing.pm_in = pm_in
                existing.pm_out = pm_out
                existing.note = note
                updated += 1
            else:
                db.add(
                    AttendanceRecord(
                        employee_id=employee.id,
                        date=normalized_date,
                        am_in=am_in,
                        am_out=am_out,
                        pm_in=pm_in,
                        pm_out=pm_out,
                        note=note,
                    )
                )
                imported += 1
        except ValueError as exc:
            skipped += 1
            errors.append(f"Row {index}: {exc}")

    db.commit()
    return {
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
    }